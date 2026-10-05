#!/usr/bin/env python3
"""Costruisce output/Dupuy_Malesia_Candidati_Raggruppati.xlsx dai risultati grezzi della ricerca
(15 agenti WebSearch, 2026-10-05). Dati NON verificati: DA VERIFICARE UMANAMENTE."""
from collections import Counter
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

SETTORI = ["Cleaning", "Utensileria", "Depolverazione", "ATEX", "Sabbiatura", "Welding", "Meccanica"]
ND = "Sede non determinata / nazionale"
C, D, M, P, K, X = ("Candidato", "Debole / da verificare", "Marketplace",
                    "Produttore / possibile concorrente", "Conflitto di marchio", "Probabile cliente finale")
# (stato, azienda, area, sito, telefono, [settori], valutazione, note)
R = [
("Johor","Corroblast","Johor Bahru","https://sgpgrid.com/company/corroblast-ID0000",None,["Sabbiatura"],C,"Attrezzature sabbiatura, rivestimenti, DPI. Fondata 1997. Sito proprio non trovato (pagina directory)."),
("Johor","Sumitec Sdn Bhd","Johor Bahru",None,None,["Utensileria"],C,"Utensili di precisione, misura, manutenzione. Fondata 2007."),
("Johor","GL Bosun Diamond Tools (M) Sdn Bhd","Muar",None,None,["Utensileria"],C,"Utensili e ferramenta, fondata 2000."),
("Johor","An Da Hardware Sdn Bhd","Johor Bahru",None,None,["Utensileria","Welding"],D,"Elettroutensili, saldatrici, DPI. Fonte poco affidabile."),
("Johor","Cleanpro Southern Sdn Bhd","Taman Kempas Utama, Johor Bahru",None,None,["Cleaning"],P,"Detergenti e chimici, non attrezzature: probabilmente fuori target."),
("Johor","Maxmedia (Odoo blog)","Johor Bahru","https://maxmedia4.odoo.com",None,["Cleaning"],D,"Chimici per pulizia industriale. Fonte debole, probabilmente fuori target."),
("Johor","Welding Alloys (Far East) Sdn Bhd","Taman Gembira, Johor Bahru",None,None,["Welding"],P,"Produttore di saldatrici automatiche e fili animati."),
("Kelantan","SLS Kelantan Sdn. Bhd.","Jalan Kuala Krai, Kota Bharu",None,"09-744 9555",["Welding","Meccanica"],D,"Attività non confermata."),
("Kelantan","KB Genset","Jalan Mahmud, Kota Bharu",None,"019-379 3191",["Meccanica"],D,"Probabile generatori, attività non confermata."),
("Kelantan","New Ah Piar Motor","Kampung Wakar Che Yeh, Kota Bharu",None,"09-748 1497",["Meccanica"],X,"Probabile officina/motori."),
("Kelantan","G&T Hyper Hardware","Sede citata a Sungai Pelong (stato non confermato)",None,None,["Utensileria"],D,"Catena ferramenta; presenza in Kelantan non confermata."),
("Kedah","H.E. Trading (S.P) Sdn. Bhd.","Sungai Petani (Jalan Kuala Ketil)",None,None,["Utensileria","Welding","Cleaning"],C,"Rivenditore dal 1991: utensili, saldatura, aspirapolveri, compressori. Candidato più rilevante in Kedah."),
("Kedah","BTS Tools Manufacturing (M) Sdn Bhd","Sungai Petani (Amanjaya)",None,"04-441 7335",["Utensileria"],P,"Il nome indica un produttore."),
("Penang","Min Hardware & Industrial Gases Sdn. Bhd.","Mak Mandin, Butterworth",None,None,["Welding"],C,"Gas industriali e attrezzature di saldatura. Dichiara di servire anche Kedah."),
("Perlis","Bann Edar Sdn Bhd","Jalan Kangar-Alor Setar, Kangar",None,None,["Meccanica"],D,"Forniture elettriche, non confermato."),
("Perlis","Siang Electrical Kangar","Kangar Town",None,None,["Meccanica"],D,"Probabile retail al dettaglio."),
("Pahang","KS Tools Depot (Malaysia)","Perkampungan Setali, Kuantan","https://mall.lapasar.com/supplier-profile/ks-tools-depot-malaysia",None,["Utensileria"],C,"Fornitore B2B wholesale (power tools, elettrico). Unico candidato locale in Pahang."),
("Negeri Sembilan","Genesis Asia Sdn Bhd","Seremban 2",None,None,["Cleaning"],X,"Impresa di servizi di pulizia: probabile cliente finale."),
("Sabah","Powermax Trading & Supplies","Inanam, Kota Kinabalu","https://mall.lapasar.com/supplier-profile/powermax-trading-supplies",None,["Utensileria"],C,"Utensili elettrici e batterie (Milwaukee). Probabile rivenditore."),
("Sabah","Summit Kinabalu Sdn Bhd","Jalan Tuaran, Kota Kinabalu",None,"+6088216946",["Meccanica"],D,"Ferramenta e vernici; fit con Dupuy debole."),
("Sabah","DoubleGain Industrial Supplies","Kolombong, Inanam, Kota Kinabalu",None,None,["Meccanica"],D,"Vetroresina, resine, pigmenti; fit debole."),
("Sabah","Unimekar Metals","KKIP, Kota Kinabalu",None,None,["Welding","Meccanica"],X,"Riparazioni e zincatura: officina/cliente finale."),
(ND,"MBJ Bina Jaya Machinery & Hardware Sdn Bhd","Sede non confermata (Sabah o Sarawak)","https://mall.lapasar.com/supplier-profile/mbj-bina-jaya-machinery-hardware",None,["Utensileria","Meccanica"],D,"Compare per Kota Kinabalu e Sarawak; sede non confermata."),
("Sarawak","Pan Sarawak Co Sdn Bhd (Pansar)","Sibu, Miri, Kuching","https://pansar.com.my",None,["Welding","Meccanica"],C,"Distributore Lincoln Electric per East Malaysia. Verificare eventuale conflitto di marca."),
("Sarawak","Sara Besi Hardware & Engineering Sdn Bhd","Jln Padungan, Kuching",None,None,["Utensileria"],C,"Materiali edili ed elettroutensili (Makita, HiKOKI)."),
("Sarawak","Sin Yew Seng Hardware & Machineries Sdn Bhd","Khoo Peng Loong Road, Sibu",None,None,["Utensileria","Meccanica"],C,"Grossista macchinari e attrezzature industriali dal 1990."),
("Sarawak","Miri Industrial Sdn Bhd","Piasau Business Park, Miri",None,None,["Meccanica"],D,"Ricambi macchine movimento terra; probabilmente fuori target."),
("Sarawak","Kiat Seng Welding Works Sdn Bhd","Piasau Industrial Estate, Miri",None,"085-656 212",["Welding"],X,"Probabile officina di saldatura."),
("Perak","Wah Seng Sin Kee Co Sdn Bhd (WSSK)","Ipoh (filiali KL, Kuantan, Butterworth)",None,None,["Utensileria","Meccanica"],C,"Distributore di ferramenta dagli anni cinquanta."),
("Perak","JESS Technology","Ipoh","https://www.jesstechnology.com",None,["Welding"],D,"Riparazione saldatrici; vendita non confermata."),
(ND,"Asia Airblast","Sede non indicata",None,None,["Sabbiatura"],D,"Sede non determinata."),
("Penang","Chiptronics (M) Sdn. Bhd.","Gelugor",None,None,["Cleaning"],D,"Attrezzature/utensili/adesivi industriali; da verificare se rivenditore."),
("Penang","Alloyplas Engineering Sdn Bhd","Penang",None,None,["Depolverazione"],P,"Dust collector, fumi di saldatura: produttore/installatore."),
("Penang","EC Pneumatic & Hardware Sdn Bhd","Jalan Chain Ferry, Perai",None,None,["Meccanica"],C,"Componenti pneumatici e hardware."),
("Penang","Berjaya Hardware (B'Worth) Sdn Bhd","Butterworth",None,None,["Meccanica","Utensileria","Welding"],C,"Hardware industriale, macchinari, utensili. Dal 1983."),
("Penang","Textran Industries (Penang) Sdn Bhd","Penang",None,None,["Meccanica"],C,"Distributore/stockist fasteners e ferramenta elettromeccanica."),
("Penang","Alloytool (M) Sdn Bhd","Perai",None,None,["Utensileria"],C,"Utensili da taglio in metallo duro, ceramica, CBN, PCD."),
("Penang","Dynarco Sdn Bhd","Perai",None,None,["Utensileria"],C,"Macchinari per precision tooling."),
("Penang","EDM-Tools (Penang) Sdn Bhd","Penang",None,None,["Utensileria"],C,"Materiali e accessori EDM, dal 1988."),
("Melaka","Tiam Heng Machinery Sdn Bhd","Melaka",None,None,["Utensileria","Meccanica"],C,"Macchinari, hardware e utensili industriali dal 1993."),
("Melaka","Yee Tat Hardware Sdn Bhd","Jonker Street, Melaka",None,None,["Utensileria"],D,"Dal 1980; probabile negozio al dettaglio."),
("Melaka","Multitech Industry Supplies","Taman Industri Malim Jaya, Melaka",None,None,["Utensileria"],C,"Strumenti di misura e collaudo di precisione."),
("Melaka","Leeden Powerweld Sdn Bhd","Air Keroh, Melaka",None,None,["Welding"],P,"Produttore di prodotti per saldatura (gruppo Leeden NOX)."),
("Melaka","iMEC Group of Companies","Melaka e Kuantan","https://imec.com.my",None,["Cleaning"],P,"Produttore di detergenti chimici."),
(ND,"UMW (Holdings / Equipment)","Non indicata","https://simeumw.com",None,["Cleaning"],K,"Rappresenta Tennant: probabile conflitto di marca."),
(ND,"Iklim Hardware & Machinery","Sede non confermata",None,None,["Utensileria","Meccanica"],D,"Dichiara 50 anni di attività; sede non chiara."),
("Selangor","Techno Tools & Equipment","Ampang","https://techno.my",None,["Cleaning","Utensileria","Welding"],C,"Fornitore autorizzato Kärcher; potrebbe essere più un retailer."),
("Selangor","CT Hardware","Petaling Jaya","https://cthardware.com",None,["Meccanica","Welding","Utensileria"],C,"Distributore ventilatori, soffianti, saldatrici."),
("Selangor","Singuan Blasting Sdn Bhd","Shah Alam",None,None,["Sabbiatura","Depolverazione"],P,"Costruisce sabbiatrici e depolveratori."),
("Selangor","Medias Industrial Supplies Sdn Bhd","Shah Alam",None,None,["Sabbiatura"],C,"Dal 1986: abrasivi, cabine, granigliatrici, ricambi."),
("Selangor","Schmidt Abrasive Blasting","Petaling Jaya",None,None,["Sabbiatura"],K,"Distributore marchio Schmidt, dal 2012: probabile conflitto di marca."),
("Selangor","See Kwong Electric (KL) Sdn Bhd","Puchong (Jalan TPP 5/2)",None,None,["ATEX"],C,"Prodotti antideflagranti ATEX/IECEx, attiva dal 1982."),
("Selangor","Ampmech Sdn Bhd","Damansara Jaya, Petaling Jaya",None,None,["ATEX"],C,"Channel partner Ceag, Rittal, Mennekes; uffici anche a Penang, Bintulu, Miri."),
(ND,"GTE (Malaysia) / GTE Lighting & Equipment","Sede non verificata","https://gte.com.my",None,["ATEX"],D,"Illuminazione antideflagrante, distributore dal 2004."),
("Selangor","Atlas Copco Malaysia","Shah Alam","https://atlascopco.com/en-my",None,["Depolverazione"],P,"Multinazionale del vuoto, non idonea come distributore."),
("Selangor","Beckhoff Automation Sdn Bhd","Kota Damansara","https://beckhoff.com",None,["ATEX"],P,"Produttore di automazione: da escludere."),
("Selangor","TSIS Welding Solution Sdn Bhd","Selangor (città n.d.)",None,None,["Welding"],K,"Distributore esclusivo Fronius dal 2015."),
("Selangor","C-MART & BUDDY TOOLS","Non chiara",None,None,["Utensileria","Welding"],D,"Probabile entità indonesiana; legame con Malesia incerto."),
("Selangor","Duromac (M) Sdn Bhd","Puchong",None,None,["Cleaning","Depolverazione"],C,"Spazzatrici; rappresenta DISAB-TELLA (aspiratori industriali)."),
("Selangor","XTS Technologies Sdn Bhd","Taman Industri Meranti Jaya, Puchong","https://www.plantautomation-technology.com/suppliers/xts-technologies-sdn-bhd",None,["Cleaning"],D,"Parts cleaning e ultrasuoni."),
("Selangor","Kärcher Cleaning Systems Sdn Bhd","Shah Alam",None,None,["Cleaning"],P,"Filiale di produttore."),
("Selangor","Klenco (M) Sdn Bhd","Bukit Naga Industrial Park",None,None,["Cleaning"],D,"Tipo di prodotto non chiaro."),
("Selangor","Tools & Machinery Parts Supplies Sdn Bhd (TMP)","Seksyen 51A, Petaling Jaya",None,None,["Utensileria","Meccanica"],C,"Fornitore industriale dal 1984."),
("Selangor","Dian Be Hardware Co Sdn Bhd","Petaling Jaya",None,None,["Utensileria"],C,"Ferramenta e utensili."),
("Selangor","Aik Huat Hardware Trading Sdn Bhd","Klang Valley",None,None,["Utensileria"],D,"Edilizia ed elettroutensili, forse poco industriale."),
("Selangor","DAV Engineering Sdn Bhd","Shah Alam","https://www.exporthub.com/dav-engineering-sdn-bhd/",None,["Depolverazione","Meccanica"],P,"Produttore di dust collector e ventilatori."),
("Selangor","HLH Welding Supply","Klang",None,None,["Welding","Utensileria"],C,"Abrasivi, utensili, gas, saldatura, compressori."),
("Selangor","Weld Systems Sdn Bhd","Petaling Jaya",None,None,["Welding"],C,"Vendita, noleggio e riparazione saldatrici e plasma."),
("Selangor","TSM Welding Technology Sdn Bhd","Klang",None,None,["Welding"],D,"Saldatura a resistenza: nicchia."),
("Selangor","MECHKINARC (M) Sdn Bhd","Shah Alam",None,None,["Welding"],C,"Macchine taglio e saldatura marchio MORROW."),
("Selangor","Systemair Malaysia","Kampung Jaya, Selangor 47000","https://www.systemair.my","603-6157 1177",["Meccanica"],P,"Probabile costruttore/filiale (ventilazione)."),
("Selangor","Elta Fans Malaysia Sdn Bhd","Shah Alam",None,None,["Meccanica"],C,"Vendita e distribuzione ventilatori."),
("Selangor","Vostermans Ventilation Sdn Bhd","Klang",None,None,["Meccanica"],P,"Probabile costruttore."),
(ND,"AMICO Bearings (M) Sdn Bhd","Rete nazionale",None,None,["Meccanica"],D,"Stockist dal 1993; filiale in Kedah non confermata."),
(ND,"Guhring Malaysia Sdn Bhd","Presenza a Penang non confermata","https://malaysia.ahk.de/en/members/members-directory/guhring-malaysia-sdn-bhd",None,["Utensileria"],P,"Filiale di produttore."),
(ND,"Lapasar (marketplace B2B MRO)","Consegna nazionale","https://lapasar.com",None,["Cleaning","Utensileria","Meccanica","Welding"],M,"Aggregatore con 10.000+ fornitori, non distributore diretto. Compare nei risultati di quasi tutti gli stati."),
]
STATI = ["Johor","Kedah","Kelantan","Melaka","Negeri Sembilan","Pahang","Penang","Perak","Perlis","Sabah","Sarawak","Selangor","Terengganu","Kuala Lumpur","Labuan","Putrajaya",ND]
rows = []
for st,az,area,sito,tel,sets,val,note in R:
    for s in sets:
        rows.append((st,s,az,area,sito,tel,val,note))
rows.sort(key=lambda r:(STATI.index(r[0]),SETTORI.index(r[1]),r[2].lower()))

wb = Workbook()
hf = PatternFill("solid", fgColor="1F3864"); hfont = Font(bold=True, color="FFFFFF")
# --- Riepilogo
ws = wb.active; ws.title = "Riepilogo"
ws.append(["Candidati distributori Dupuy - Malesia (righe stato x settore). DATI GREZZI: DA VERIFICARE UMANAMENTE"])
ws["A1"].font = Font(bold=True, size=12)
ws.append(["Totale = tutte le valutazioni; 'Candidati' = solo valutazione Candidato"])
ws.append([])
ws.append(["Stato"] + SETTORI + ["Aziende uniche", "Di cui Candidato"])
for c in ws[4]: c.fill = hf; c.font = hfont
cnt = Counter((r[0], r[1]) for r in rows)
for st in STATI:
    uniq = {r[2] for r in rows if r[0]==st}
    cand = {r[2] for r in rows if r[0]==st and r[6]==C}
    ws.append([st] + [cnt.get((st,s),0) or None for s in SETTORI] + [len(uniq), len(cand)])
ws.column_dimensions["A"].width = 34
for i in range(2, 11): ws.column_dimensions[get_column_letter(i)].width = 15
ws.append([]); ws.append(["Terengganu, Kuala Lumpur, Labuan e Putrajaya: nessun candidato locale trovato (i risultati Klang Valley sono assegnati a Selangor)."])
# --- Tutti
wt = wb.create_sheet("Tutti i risultati")
head = ["stato","settore","ragione_sociale","area","sito_web","telefono","email","valutazione","note","stato_verifica"]
wt.append(head)
for c in wt[1]: c.fill = hf; c.font = hfont
for r in rows:
    wt.append([r[0],r[1],r[2],r[3],r[4],r[5],None,r[6],r[7],"DA VERIFICARE UMANAMENTE"])
for col,w in zip("ABCDEFGHIJ",[26,16,42,38,40,16,18,30,70,26]): wt.column_dimensions[col].width = w
wt.freeze_panes = "C2"; wt.auto_filter.ref = wt.dimensions
for row in wt.iter_rows(min_row=2):
    for c in row: c.alignment = Alignment(wrap_text=True, vertical="top")
# --- Un foglio per stato
for st in STATI:
    sub = [r for r in rows if r[0]==st]
    if not sub: continue
    w = wb.create_sheet(st[:31] if st!=ND else "Sede n.d.")
    w.append(head[1:])
    for c in w[1]: c.fill = hf; c.font = hfont
    for r in sub: w.append([r[1],r[2],r[3],r[4],r[5],None,r[6],r[7],"DA VERIFICARE UMANAMENTE"])
    for col,wd in zip("ABCDEFGHI",[16,42,38,40,16,18,30,70,26]): w.column_dimensions[col].width = wd
    w.freeze_panes = "A2"
OUT = "output/Dupuy_Malesia_Candidati_Raggruppati.xlsx"
wb.save(OUT)
print(len(rows), "righe;", len({r[2] for r in rows}), "aziende uniche ->", OUT)
